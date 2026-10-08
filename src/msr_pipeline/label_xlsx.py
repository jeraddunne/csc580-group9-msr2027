"""Excel labelling workbook for the P-01 validation kit.

`python scripts/annotation_kit.py xlsx --rater jd --round 1` writes
data/annotations/work/labelling_jd_r1.xlsx (gitignored, because it quotes dataset text). It is
the same task as the labelling studio, as an answer sheet: one row per label row, dropdowns for
the guideline's labels, the guideline itself, live completeness checks, and a Metrics sheet
whose formulas use the same definitions as `annotation_kit.py score`.

`python scripts/annotation_kit.py import data/annotations/work/labelling_jd_r1.xlsx` reads it
back through the same validation as a CSV from the studio.

The workbook never chooses or fills in a label. Its input cells start from the rater's own
label files and work log, and are empty for a rater who has not labelled yet.
"""

from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path

import pandas as pd

from . import label_ui
from . import validation as v

FORMAT = "p01-labelling-workbook/1"
EXCERPT_CHARS = 1200
SHEET = {"signals": "Signals", "lineage": "Lineage", "drift": "Drift"}
FONT = "Arial"

# Fills: inputs the rater fills, calculated cells, headers.
INPUT_FILL = "FFF2CC"
CALC_FILL = "F2F2F2"
HEAD_FILL = "1F4E79"
DONE_FILL = "C6EFCE"
BAD_FILL = "FFC7CE"

# (key, header, width, role). Roles: info (filled in, read only), key (identifies the row for
# the import), input (the rater fills it), calc (formula), helper (hidden formula).
SIGNAL_COLS = [
    ("row", "Row", 6, "info"),
    ("item", "Item", 6, "info"),
    ("stratum", "Drawn for", 15, "info"),
    ("rule_id", "Rule", 11, "key"),
    ("category", "Category", 17, "info"),
    ("severity", "Severity", 8, "info"),
    ("question", "What to decide", 34, "info"),
    ("excerpt", "Matched text (read only: never run, open, or search for it)", 60, "info"),
    ("packet", "Reading packet", 12, "info"),
    ("label", "Label", 17, "input"),
    ("missed_category", "Missed category", 18, "input"),
    ("notes", "Notes and reason", 46, "input"),
    ("why", "What settled it", 34, "input"),
    ("confidence", "Confidence", 11, "input"),
    ("date", "Date labelled", 12, "input"),
    ("minutes", "Minutes", 8, "input"),
    ("reason_needed", "Reason needed", 9, "calc"),
    ("status", "Status", 15, "calc"),
    ("round2", "Round 2 from", 12, "calc"),
    ("file_sha", "file_sha", 44, "key"),
    ("h_rank", "rank", 6, "helper"),
    ("h_first", "first of skill and category", 6, "helper"),
    ("h_cat", "skill label for category", 6, "helper"),
]
LINEAGE_COLS = [
    ("row", "Pair", 6, "info"),
    ("source", "Drawn as", 10, "info"),
    ("similarity", "Similarity", 9, "info"),
    ("names", "Front-matter names", 30, "info"),
    ("excerpt", "Diff, a to b (read only)", 60, "info"),
    ("packet", "Reading packet", 12, "info"),
    ("label", "Label", 14, "input"),
    ("notes", "Notes and reason", 46, "input"),
    ("confidence", "Confidence", 11, "input"),
    ("date", "Date labelled", 12, "input"),
    ("minutes", "Minutes", 8, "input"),
    ("reason_needed", "Reason needed", 9, "calc"),
    ("status", "Status", 15, "calc"),
    ("round2", "Round 2 from", 12, "calc"),
    ("file_sha_a", "file_sha_a", 44, "key"),
    ("file_sha_b", "file_sha_b", 44, "key"),
]
DRIFT_COLS = [
    ("row", "Pair", 6, "info"),
    ("similarity", "Similarity", 9, "info"),
    ("ordered", "Ordered (a older)", 9, "info"),
    ("only_in_a", "High-risk only in a", 18, "info"),
    ("only_in_b", "High-risk only in b", 18, "info"),
    ("names", "Front-matter names", 30, "info"),
    ("excerpt", "Diff, a to b (read only)", 60, "info"),
    ("packet", "Reading packet", 12, "info"),
    ("label", "Label", 20, "input"),
    ("notes", "Notes and reason", 46, "input"),
    ("confidence", "Confidence", 11, "input"),
    ("date", "Date labelled", 12, "input"),
    ("minutes", "Minutes", 8, "input"),
    ("reason_needed", "Reason needed", 9, "calc"),
    ("status", "Status", 15, "calc"),
    ("round2", "Round 2 from", 12, "calc"),
    ("file_sha_a", "file_sha_a", 44, "key"),
    ("file_sha_b", "file_sha_b", 44, "key"),
]
COLUMNS = {"signals": SIGNAL_COLS, "lineage": LINEAGE_COLS, "drift": DRIFT_COLS}

_FIELD = re.compile(r"^- ([a-z_ -]+): (.*)$", re.M)


def workbook_filename(rater: str, round_: int) -> str:
    v.label_filename("signals", rater, round_)  # same rater and round checks as label files
    return f"labelling_{rater}_r{round_}.xlsx"


def _require_openpyxl():
    try:
        import openpyxl  # noqa: F401
    except ImportError as exc:  # pragma: no cover - exercised only without the dependency
        raise RuntimeError(
            "The Excel workbook needs openpyxl: pip install -r requirements.txt"
        ) from exc


# ---------------------------------------------------------------------------
# Reading packets
# ---------------------------------------------------------------------------


def packet_files(packet_dir: Path, kind: str) -> dict[str, str]:
    """Item key -> packet file name, for links from the workbook to the full reading packet."""
    out: dict[str, str] = {}
    if not packet_dir.exists():
        return out
    for path in sorted(packet_dir.glob("*.md")):
        if path.name == "INDEX.md":
            continue
        key = label_ui.packet_key(kind, path.read_text(encoding="utf-8"))
        if key:
            out.setdefault(key, path.name)
    return out


def packet_fields(packet: str) -> dict[str, str]:
    return {m.group(1): m.group(2).replace("`", "").strip() for m in _FIELD.finditer(packet)}


def _dedent(lines: list[str]) -> str:
    out = []
    for line in lines:
        if line.startswith("Rule note:"):
            continue
        out.append(line[4:] if line.startswith("    ") else line)
    text = "\n".join(out).strip()
    return re.sub(r"\n{3,}", "\n\n", text)


def _clip(text: str) -> str:
    if len(text) <= EXCERPT_CHARS:
        return text
    return text[:EXCERPT_CHARS].rstrip() + "\n... (cut here; open the reading packet for the rest)"


def excerpt(kind: str, packet: str, rule_id: str = "") -> str:
    """The part of a reading packet that a row is about, as plain text."""
    lines = packet.splitlines()
    if kind == "signals" and rule_id and rule_id != v.NONE_RULE:
        start = next((i for i, x in enumerate(lines) if x.startswith(f"## {rule_id} (")), None)
        if start is None:
            return ""
        end = next(
            (i for i in range(start + 1, len(lines)) if lines[i].startswith("## ")), len(lines)
        )
        return _clip(_dedent(lines[start + 1 : end]))
    marker = "No rule matched" if kind == "signals" else "Unified diff"
    start = next((i for i, x in enumerate(lines) if x.startswith(marker)), None)
    if start is None:
        return ""
    return _clip(_dedent(lines[start + 1 :]))


# ---------------------------------------------------------------------------
# Cell helpers
# ---------------------------------------------------------------------------


def _styles():
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    thin = Side(style="thin", color="D9D9D9")
    return {
        "font": Font(name=FONT, size=10),
        "bold": Font(name=FONT, size=10, bold=True),
        "head": Font(name=FONT, size=10, bold=True, color="FFFFFF"),
        "title": Font(name=FONT, size=14, bold=True),
        "h2": Font(name=FONT, size=12, bold=True),
        "muted": Font(name=FONT, size=9, italic=True, color="595959"),
        "link": Font(name=FONT, size=10, color="0563C1", underline="single"),
        "input": PatternFill("solid", fgColor=INPUT_FILL),
        "calc": PatternFill("solid", fgColor=CALC_FILL),
        "headfill": PatternFill("solid", fgColor=HEAD_FILL),
        "wrap": Alignment(wrap_text=True, vertical="top"),
        "top": Alignment(vertical="top"),
        "border": Border(bottom=thin),
    }


def _text(cell, value, font) -> None:
    """Write text literally: illegal characters dropped, never read as a formula."""
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    s = ILLEGAL_CHARACTERS_RE.sub("", "" if value is None else str(value))
    cell.value = s if s else None
    if s.startswith("="):
        cell.data_type = "s"
    cell.font = font


def _write_lines(ws, row: int, items: list, st: dict, col: int = 1) -> int:
    """Write (style, text) pairs down a column, one per row; returns the next free row."""
    for style, text in items:
        cell = ws.cell(row=row, column=col)
        _text(cell, text, st[style])
        if style in ("font", "muted"):
            cell.alignment = st["wrap"]
        row += 1
    return row


# ---------------------------------------------------------------------------
# Building
# ---------------------------------------------------------------------------


def _rule_index(rule_file: dict) -> dict[str, dict]:
    return {
        str(r["id"]): r
        for r in rule_file.get("rules") or []
        if str(r.get("status", "active")) == "active"
    }


def _worklog_index(worklog: pd.DataFrame | None) -> dict[tuple, dict]:
    if worklog is None or worklog.empty:
        return {}
    out = {}
    for rec in worklog.fillna("").astype(str).to_dict("records"):
        out[(rec["key_a"], rec["key_b"], rec["rule_id"])] = rec
    return out


def _guideline_minutes_per_row(kind: str, items: list[dict], guide: dict | None) -> float:
    times = (guide or {}).get("times", {})

    def pick(prefix: str, default: str) -> int:
        return label_ui._minutes(
            next((t for k, t in times.items() if k.startswith(prefix)), default)
        )

    rows = sum(len(i["rows"]) for i in items) or 1
    if kind == "signals":
        total = sum(
            pick("No-signal", "4")
            if any(r["rule_id"] == v.NONE_RULE for r in i["rows"])
            else pick("Signal", "3")
            for i in items
        )
    else:
        total = len(items) * pick(
            "Lineage" if kind == "lineage" else "Drift", "2" if kind == "lineage" else "4"
        )
    return round(total / rows, 2)


def build_workbook(
    out: Path,
    rater: str,
    round_: int,
    kinds: dict[str, dict],
    guide: dict | None,
    rule_file: dict,
    sample_info: dict | None = None,
) -> dict[str, int]:
    """Write the workbook. ``kinds`` maps kind -> {items, packet_files, packet_dir, worklog}.

    ``items`` come from label_ui.build_items (labels from the rater's own file); packet links
    are relative to the workbook's folder. Returns rows written per kind.
    """
    _require_openpyxl()
    from openpyxl import Workbook
    from openpyxl.workbook.defined_name import DefinedName

    st = _styles()
    wb = Workbook()
    wb.remove(wb.active)
    rules = _rule_index(rule_file)
    categories = list((rule_file.get("categories") or {}).keys())
    reasons = label_ui.reason_types(guide)

    start = wb.create_sheet("Start here")
    howto = wb.create_sheet("How to decide")
    sheets = {k: wb.create_sheet(SHEET[k]) for k in v.KINDS if k in kinds}
    metrics = wb.create_sheet("Metrics")
    amap = wb.create_sheet("Assignment map")
    rules_ws = wb.create_sheet("Rules")
    lists = wb.create_sheet("Lists")
    meta = wb.create_sheet("_meta")

    list_ranges = _write_lists(lists, st, categories, reasons)
    counts: dict[str, int] = {}
    ranges: dict[str, dict] = {}
    for kind, ws in sheets.items():
        info = kinds[kind]
        counts[kind], ranges[kind] = _write_kind_sheet(
            ws, kind, rater, round_, info, rules, reasons, list_ranges, st
        )
    _write_start(start, rater, round_, kinds, guide, reasons, st)
    _write_howto(howto, guide, st)
    _write_rules(rules_ws, rule_file, st)
    _write_metrics(
        metrics, kinds, ranges, rules, categories, reasons, guide, sample_info, rater, round_, st
    )
    _write_assignment_map(amap, [k for k in v.KINDS if k in kinds], st)

    meta_rows = [
        ("format", FORMAT),
        ("rater", rater),
        ("round", str(round_)),
        ("kinds", ",".join(k for k in v.KINDS if k in kinds)),
        ("built_at", datetime.now().isoformat(timespec="seconds")),
        ("guideline_version", (guide or {}).get("version", "")),
    ]
    for i, (k, val) in enumerate(meta_rows, start=1):
        _text(meta.cell(row=i, column=1), k, st["font"])
        _text(meta.cell(row=i, column=2), val, st["font"])
    meta.sheet_state = "hidden"
    lists.sheet_state = "hidden"
    wb.defined_names["Z95"] = DefinedName("Z95", attr_text="Metrics!$B$5")
    # openpyxl stores formulas without results; ask Excel to calculate them on opening.
    wb.calculation.fullCalcOnLoad = True
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return counts


def _write_lists(ws, st, categories, reasons) -> dict[str, str]:
    columns = {
        "rule_labels": list(v.SIGNAL_LABELS),
        "none_labels": list(v.NO_SIGNAL_LABELS),
        "categories": categories,
        "reasons": [r["text"] for r in reasons],
        "confidence": list(label_ui.CONFIDENCE),
        "lineage": list(v.LINEAGE_LABELS),
        "drift": list(v.DRIFT_LABELS),
    }
    out = {}
    for c, (name, values) in enumerate(columns.items(), start=1):
        _text(ws.cell(row=1, column=c), name, st["bold"])
        for r, val in enumerate(values, start=2):
            _text(ws.cell(row=r, column=c), val, st["font"])
        letter = ws.cell(row=1, column=c).column_letter
        out[name] = f"Lists!${letter}$2:${letter}${max(2, len(values) + 1)}"
    return out


def _col_letters(cols) -> dict[str, str]:
    from openpyxl.utils import get_column_letter

    return {key: get_column_letter(i) for i, (key, *_rest) in enumerate(cols, start=1)}


def _write_kind_sheet(ws, kind, rater, round_, info, rules, reasons, list_ranges, st):
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import PatternFill, Protection
    from openpyxl.worksheet.datavalidation import DataValidation

    cols = COLUMNS[kind]
    L = _col_letters(cols)
    items = info["items"]
    files = info.get("packet_files") or {}
    packet_dir = info.get("packet_dir", "")
    log = _worklog_index(info.get("worklog"))
    reason_text = {r["id"]: r["text"] for r in reasons}
    justify = list(v.JUSTIFY_LABELS[kind])
    gap = label_ui.ROUND2_MIN_DAYS[kind]

    for c, (key, header, width, role) in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=c)
        _text(cell, header, st["head"])
        cell.fill = st["headfill"]
        cell.alignment = st["wrap"]
        ws.column_dimensions[L[key]].width = width
        if role == "helper":
            ws.column_dimensions[L[key]].hidden = True
    ws.row_dimensions[1].height = 42

    rows = []
    for n, item in enumerate(items, start=1):
        fields = packet_fields(item["packet"])
        for r in item["rows"]:
            rows.append((n, item, fields, r))
    last = len(rows) + 1
    rng = lambda key: f"${L[key]}$2:${L[key]}${last}"  # noqa: E731

    dv = {
        "rule": DataValidation(
            type="list", formula1="=" + list_ranges["rule_labels"], allow_blank=True
        ),
        "none": DataValidation(
            type="list", formula1="=" + list_ranges["none_labels"], allow_blank=True
        ),
        "pair": DataValidation(
            type="list",
            formula1="=" + list_ranges["lineage" if kind == "lineage" else "drift"],
            allow_blank=True,
        ),
        "cat": DataValidation(
            type="list", formula1="=" + list_ranges["categories"], allow_blank=True
        ),
        "why": DataValidation(type="list", formula1="=" + list_ranges["reasons"], allow_blank=True),
        "conf": DataValidation(
            type="list", formula1="=" + list_ranges["confidence"], allow_blank=True
        ),
        "date": DataValidation(
            type="date", operator="greaterThan", formula1="DATE(2026,1,1)", allow_blank=True
        ),
        "min": DataValidation(
            type="decimal", operator="between", formula1="0", formula2="600", allow_blank=True
        ),
    }
    for d in dv.values():
        d.error = "Choose a value from the list (the guideline's exact labels)."
        d.errorTitle = "Not a value this sheet accepts"
        d.showErrorMessage = True
    dv["date"].error = "Enter a date such as 2026-10-05."
    dv["min"].error = "Enter minutes as a number from 0 to 600."

    for i, (n, item, fields, r) in enumerate(rows, start=2):
        is_none = kind == "signals" and r["rule_id"] == v.NONE_RULE
        if kind == "signals":
            rule = rules.get(r["rule_id"], {})
            key_a, key_b, rid = item["key"], "", r["rule_id"]
            values = {
                "row": i - 1,
                "item": n,
                "stratum": fields.get("stratum", ""),
                "rule_id": r["rule_id"],
                "category": "" if is_none else str(rule.get("category", "")),
                "severity": "" if is_none else rule.get("severity", ""),
                "question": (
                    "No rule matched. Any high-risk capability? NONE_PRESENT, or MISSED_RISKY plus the category"
                    if is_none
                    else f"Is {rule.get('category', 'this capability')} really there as something the agent would do? RISKY, BENIGN_CONTEXT, or NOT_PRESENT"
                ),
                "excerpt": excerpt("signals", item["packet"], r["rule_id"]),
                "file_sha": item["key"],
                "label": r.get("label", ""),
                "missed_category": r.get("missed_category", ""),
                "notes": r.get("notes", ""),
            }
        else:
            key_a, key_b, rid = r["file_sha_a"], r["file_sha_b"], ""
            values = {
                "row": n,
                "source": fields.get("source", ""),
                "similarity": fields.get("similarity", ""),
                "ordered": fields.get("ordered", ""),
                "only_in_a": fields.get("only_in_a", ""),
                "only_in_b": fields.get("only_in_b", ""),
                "names": fields.get("front-matter names", ""),
                "excerpt": excerpt(kind, item["packet"]),
                "file_sha_a": r["file_sha_a"],
                "file_sha_b": r["file_sha_b"],
                "label": r.get("label", ""),
                "notes": r.get("notes", ""),
            }
        lg = log.get((key_a, key_b, rid), {})
        values["confidence"] = (
            lg.get("confidence", "") if lg.get("confidence", "") in label_ui.CONFIDENCE else ""
        )
        values["why"] = reason_text.get(lg.get("reason_type", ""), "")
        first = (lg.get("first_labelled_at") or "")[:10]
        values["date"] = (
            date.fromisoformat(first) if re.match(r"^\d{4}-\d{2}-\d{2}$", first) else None
        )
        secs = pd.to_numeric(lg.get("active_seconds", ""), errors="coerce")
        values["minutes"] = round(float(secs) / 60, 1) if pd.notna(secs) and secs > 0 else None

        lab = f"{L['label']}{i}"
        allowed = (
            list(v.NO_SIGNAL_LABELS if is_none else v.SIGNAL_LABELS)
            if kind == "signals"
            else list(v.LINEAGE_LABELS if kind == "lineage" else v.DRIFT_LABELS)
        )
        valid = "OR(" + ",".join(f'EXACT(TRIM({lab}),"{a}")' for a in allowed) + ")"
        needs = "OR(" + ",".join(f'TRIM({lab})="{a}"' for a in justify) + ")"
        status = f'IF(TRIM({lab})="","to do",IF(NOT({valid}),"invalid label",'
        if kind == "signals":
            status += f'IF(AND(TRIM({lab})="MISSED_RISKY",TRIM({L["missed_category"]}{i})=""),"needs a category",'
        status += f'IF(AND({L["reason_needed"]}{i}="yes",LEN(TRIM({L["notes"]}{i}))<{v.MIN_JUSTIFICATION_CHARS}),"needs a reason","done")'
        status += ")" + (")" if kind == "signals" else "") + ")"
        formulas = {
            "reason_needed": f'=IF({needs},"yes","")',
            "status": "=" + status,
            "round2": f'=IF(AND({L["status"]}{i}="done",ISNUMBER({L["date"]}{i})),{L["date"]}{i}+{gap},"")',
        }
        if kind == "signals":
            S = L["status"]
            formulas["h_rank"] = (
                '=""'
                if is_none
                else f'=IF({S}{i}<>"done","",IF(TRIM({lab})="RISKY",2,IF(TRIM({lab})="BENIGN_CONTEXT",1,0)))'
            )
            sha, cat = L["file_sha"], L["category"]
            formulas["h_first"] = (
                "=0"
                if is_none
                else f"=IF(COUNTIFS(${sha}$2:{sha}{i},{sha}{i},${cat}$2:{cat}{i},{cat}{i})=1,1,0)"
            )
            grp = f"{rng('file_sha')},{sha}{i},{rng('category')},{cat}{i},{rng('h_rank')}"
            formulas["h_cat"] = (
                '=""'
                if is_none
                else f'=IF({L["h_first"]}{i}<>1,"",IF(COUNTIFS({grp},2)>0,2,IF(COUNTIFS({grp},1)>0,1,IF(COUNTIFS({grp},0)>0,0,""))))'
            )

        for c, (key, _header, _w, role) in enumerate(cols, start=1):
            cell = ws.cell(row=i, column=c)
            if role == "calc" or role == "helper":
                cell.value = formulas[key]
                cell.font = st["font"]
                cell.fill = st["calc"]
                if key == "round2":
                    cell.number_format = "yyyy-mm-dd"
            elif key == "packet":
                name = files.get(item["key"], "")
                if name:
                    _text(cell, "open", st["link"])
                    cell.hyperlink = f"{packet_dir}/{name}"
                else:
                    _text(cell, "not found", st["muted"])
            elif key in ("date", "minutes"):
                cell.value = values.get(key)
                cell.font = st["font"]
                if key == "date":
                    cell.number_format = "yyyy-mm-dd"
            elif key in ("row", "item", "severity"):
                cell.value = values.get(key) if values.get(key) != "" else None
                cell.font = st["font"]
            elif key == "similarity":
                num = pd.to_numeric(values.get(key, ""), errors="coerce")
                cell.value = float(num) if pd.notna(num) else values.get(key, "")
                cell.font = st["font"]
            else:
                _text(cell, values.get(key, ""), st["font"])
            cell.alignment = st["top"]
            if role == "input":
                cell.fill = st["input"]
                cell.protection = Protection(locked=False)
                if key == "notes":
                    cell.alignment = st["wrap"]
            if key == "excerpt":
                cell.alignment = st["top"]

        if kind == "signals":
            dv["none" if is_none else "rule"].add(lab)
            if is_none:
                dv["cat"].add(f"{L['missed_category']}{i}")
            else:
                dv["why"].add(f"{L['why']}{i}")
        else:
            dv["pair"].add(lab)
        dv["conf"].add(f"{L['confidence']}{i}")
        dv["date"].add(f"{L['date']}{i}")
        dv["min"].add(f"{L['minutes']}{i}")

    for d in dv.values():
        if d.sqref.ranges:  # a kind with no rows has nothing to validate
            ws.add_data_validation(d)
    status_rng = f"{L['status']}2:{L['status']}{max(last, 2)}"
    ws.conditional_formatting.add(
        status_rng,
        CellIsRule(
            operator="equal", formula=['"done"'], fill=PatternFill("solid", fgColor=DONE_FILL)
        ),
    )
    for word in ("needs a reason", "needs a category", "invalid label"):
        ws.conditional_formatting.add(
            status_rng,
            CellIsRule(
                operator="equal", formula=[f'"{word}"'], fill=PatternFill("solid", fgColor=BAD_FILL)
            ),
        )
    ws.freeze_panes = f"{L['label']}2"
    ws.auto_filter.ref = f"A1:{L[cols[-1][0]]}{last}"
    ws.protection.sheet = True
    ws.protection.autoFilter = False
    ws.protection.sort = False
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    return len(rows), {"L": L, "last": last, "sheet": SHEET[kind], "rng": {k: rng(k) for k in L}}


def _write_start(ws, rater, round_, kinds, guide, reasons, st) -> None:
    ws.column_dimensions["A"].width = 110
    kind_list = [k for k in v.KINDS if k in kinds]
    page = ", ".join(f"label_{k}_{rater}_r{round_}.html" for k in kind_list)
    sched = []
    for k in kind_list:
        due = label_ui.ROUND_DUE.get((k, round_), "")
        if due:
            sched.append(("font", f"{SHEET[k]}: round {round_} due {due}."))
    if round_ == 1 and rater == "jd":
        sched.append(
            (
                "font",
                "Round 2 relabels a 30% subset, each item at least 14 days (signals, lineage) or 7 days "
                "(drift, optional) after its round 1 label. Every day you finish earlier gives round 2 a "
                f"day more before scoring on {label_ui.SCORE_DUE}. The Round 2 from column shows each "
                "row's earliest date once you fill in Date labelled.",
            )
        )
    sched.append(("font", f"Scoring with annotation_kit.py score is due {label_ui.SCORE_DUE}."))
    independence = []
    for k in kind_list:
        independence += [
            x for x in label_ui._independence(k, rater, round_) if x not in independence
        ]
    safety = (guide or {}).get("safety") or []
    path = f"data/annotations/work/{workbook_filename(rater, round_)}"
    lines = [
        ("title", f"P-01 labelling workbook: rater {rater}, round {round_}"),
        (
            "muted",
            f"Built {date.today().isoformat()} from docs/validation/ANNOTATION_GUIDELINE.md version "
            f"{(guide or {}).get('version', '?')} and rules/skill_risk_rules.yaml. Labels in this workbook come only "
            "from your own label files: nothing is filled in for you, and every label is your judgement.",
        ),
        ("font", ""),
        ("h2", "How to use this workbook"),
        (
            "font",
            "1. Read the How to decide sheet once, fully (about 10 minutes), and keep it open while you label.",
        ),
        (
            "font",
            "2. Work one sheet at a time: " + ", then ".join(SHEET[k] for k in kind_list) + ".",
        ),
        (
            "font",
            "3. Read the matched text in each row. If it is cut short, click open in the Reading packet "
            f"column for the whole text, or use the labelling page ({page}), which shows the same items "
            "with the matches highlighted.",
        ),
        (
            "font",
            "4. Choose the Label from the dropdown. Yellow cells are yours to fill; everything else is filled in or calculated.",
        ),
        (
            "font",
            f"5. If Reason needed says yes, write at least {v.MIN_JUSTIFICATION_CHARS} characters in Notes "
            "saying what the text actually shows. Status turns green (done) when the row is complete.",
        ),
        (
            "font",
            "6. Optional, and worth it: What settled it (false positives only), Confidence, Date labelled, "
            "and Minutes. They feed the error analysis, the round 2 dates, and your hours.",
        ),
        ("font", "7. Save often (Ctrl+S). Take a break every 30 to 40 minutes."),
        (
            "font",
            f"8. At the end of every session, import: python scripts/annotation_kit.py import {path}. "
            "It checks every label and reason and saves them into data/annotations/. Fix anything it lists, save, and import again.",
        ),
        (
            "font",
            "9. The Metrics sheet shows progress and pace. Its results unlock when every row of a kind is done.",
        ),
        ("font", ""),
        ("h2", "The workbook or the labelling page, one at a time"),
        (
            "font",
            "Use one of them for a kind at a time. To switch, import first. Then either rebuild this "
            "workbook from your imported labels (annotation_kit.py xlsx --rater "
            f"{rater} --round {round_} --force), or open the page and click Load CSV with data/annotations/<kind>_{rater}_r{round_}.csv. "
            "The import refuses to overwrite a saved label with a different one unless you add --replace.",
        ),
        ("font", ""),
        ("h2", "Colour key"),
        (
            "font",
            "Yellow: you fill it in. Grey: calculated, do not edit. Green status: done. Red status: needs attention.",
        ),
        (
            "font",
            "Sheets are protected against accidental edits to the keys and formulas (Review, Unprotect Sheet if you must; there is no password).",
        ),
        ("font", ""),
        ("h2", "Example row (invented; not one of your items)"),
    ]
    row = _write_lines(ws, 1, lines, st)
    example = [
        ("Label", "NOT_PRESENT"),
        (
            "Notes and reason",
            "matched `sudo` inside the identifier SUDO_USER in a prose sentence, not an instruction",
        ),
        ("What settled it", reasons[0]["text"] if reasons else ""),
        ("Confidence", "Sure"),
        ("Date labelled", "2026-10-05"),
        ("Minutes", "3"),
    ]
    for k, val in example:
        cell = ws.cell(row=row, column=1)
        _text(cell, f"{k}: {val}", st["font"])
        cell.fill = st["input"]
        row += 1
    tail = [("font", ""), ("h2", "Safety rules (guideline section 1)")]
    tail += [("font", f"{i}. {s}") for i, s in enumerate(safety, start=1)]
    if independence:
        tail += [("font", ""), ("h2", "Working independently")] + [
            ("font", s) for s in independence
        ]
    tail += [("font", ""), ("h2", "Schedule")] + sched
    tail += [
        ("font", ""),
        ("h2", "Keeping your work safe"),
        (
            "font",
            "This file holds quotes from the dataset, so it lives in data/annotations/work/, which git ignores. Never commit it or share it.",
        ),
        (
            "font",
            "Only the label files that the import writes (data/annotations/<kind>_<rater>_r<round>.csv) are committed, following the blindness rule in docs/validation/README.md.",
        ),
        ("font", ""),
        ("h2", "What your labels are for"),
        (
            "font",
            "See the Assignment map sheet: each number these labels produce, what it answers, and where it goes in the report and the rubric.",
        ),
    ]
    _write_lines(ws, row, tail, st)


def _write_table(ws, row, head, body, st, widths=None) -> int:
    for c, h in enumerate(head, start=1):
        cell = ws.cell(row=row, column=c)
        _text(cell, h, st["head"])
        cell.fill = st["headfill"]
        cell.alignment = st["wrap"]
    row += 1
    for rec in body:
        for c, val in enumerate(rec, start=1):
            cell = ws.cell(row=row, column=c)
            _text(cell, val, st["font"])
            cell.alignment = st["wrap"]
        row += 1
    if widths:
        from openpyxl.utils import get_column_letter

        for c, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(c)].width = max(
                ws.column_dimensions[get_column_letter(c)].width or 0, w
            )
    return row + 1


def _write_howto(ws, guide, st) -> None:
    from openpyxl.utils import get_column_letter

    for c, w in enumerate([24, 60, 50, 50], start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    if not guide:
        _text(
            ws.cell(row=1, column=1),
            "The guideline was not found when this workbook was built. Read docs/validation/ANNOTATION_GUIDELINE.md.",
            st["bold"],
        )
        return
    g = guide
    row = _write_lines(
        ws,
        1,
        [
            (
                "title",
                f"How to decide (ANNOTATION_GUIDELINE.md version {g['version']}, {g['date']})",
            ),
            (
                "muted",
                "Copied from the guideline when this workbook was built. The guideline file is the authority.",
            ),
            ("font", ""),
            ("h2", "Signals: the question for a rule row"),
            ("font", g["signals"]["question"]),
            (
                "h2",
                "Decision order (guideline 2.2): answer in order; the first one that settles it gives the label",
            ),
        ],
        st,
    )
    steps = []
    for i, s in enumerate(g["signals"]["decision"], start=1):
        settles = (
            ""
            if i == len(g["signals"]["decision"])
            else ("on no" if s["label"] == "NOT_PRESENT" else "on yes")
        )
        steps.append([f"Q{i}", s["question"], s["detail"], f"{s['label']} {settles}".strip()])
    row = _write_table(ws, row, ["Question", "Ask", "Detail", "Settles as"], steps, st)
    row = _write_table(
        ws,
        row,
        ["Signal label", "Meaning"],
        [[k, d] for k, d in g["signals"]["labels"].items()],
        st,
    )
    row = _write_table(
        ws,
        row,
        ["No-signal label", "Meaning"],
        [[k, d] for k, d in g["signals"]["none"].items()],
        st,
    )
    if g["signals"].get("noneNote"):
        row = _write_lines(ws, row, [("font", g["signals"]["noneNote"]), ("font", "")], st)
    examples = [
        [
            ex["capability"],
            ex.get("RISKY", ""),
            ex.get("BENIGN_CONTEXT", ""),
            ex.get("NOT_PRESENT", ""),
        ]
        for ex in {e["capability"]: e for e in g["signals"]["examples"].values()}.values()
    ]
    row = _write_lines(
        ws, row, [("h2", "Examples by capability (guideline 2.3, all invented)")], st
    )
    row = _write_table(
        ws, row, ["Capability", "RISKY", "BENIGN_CONTEXT", "NOT_PRESENT"], examples, st
    )
    row = _write_lines(
        ws,
        row,
        [("h2", "Special cases (guideline 2.4)")]
        + [("font", "- " + s) for s in g["signals"]["special"]]
        + [("font", "")],
        st,
    )
    for kind in ("lineage", "drift"):
        k = g[kind]
        row = _write_lines(
            ws,
            row,
            [("h2", f"{kind.capitalize()}: {k['question']}")]
            + ([("font", k["intro"])] if k.get("intro") else []),
            st,
        )
        row = _write_table(
            ws,
            row,
            ["Label", "Meaning", "Example"],
            [[lab, d.get("meaning", ""), d.get("example", "")] for lab, d in k["labels"].items()],
            st,
        )
        if k.get("note"):
            row = _write_lines(ws, row, [("font", k["note"]), ("font", "")], st)
    row = _write_lines(ws, row, [("h2", "Reasons (guideline 5)")], st)
    row = _write_table(
        ws,
        row,
        ["Kind", "Labels that need a reason of at least 15 characters"],
        [[kind, ", ".join(v.JUSTIFY_LABELS[kind])] for kind in v.KINDS],
        st,
    )
    row = _write_lines(ws, row, [("font", g["reason"]), ("font", "")], st)
    _write_table(
        ws, row, ["Item", "Time (guideline 7)"], [[k, t] for k, t in g["times"].items()], st
    )


def _write_rules(ws, rule_file, st) -> None:
    cats = rule_file.get("categories") or {}
    body = []
    for r in rule_file.get("rules") or []:
        if str(r.get("status", "active")) != "active":
            continue
        ex = r.get("examples") or {}
        sev = int(r["severity"])
        body.append(
            [
                str(r["id"]),
                str(r["category"]),
                str(sev),
                "yes" if sev >= label_ui.HIGH_RISK_SEVERITY else "no, context only",
                str(r.get("note", "")),
                "; ".join(str(x) for x in ex.get("match") or []),
                "; ".join(str(x) for x in ex.get("no_match") or []),
                str(cats.get(r["category"], "")),
            ]
        )
    row = _write_lines(
        ws, 1, [("title", "Rules (rules/skill_risk_rules.yaml, active rules)"), ("font", "")], st
    )
    _write_table(
        ws,
        row,
        [
            "Rule",
            "Category",
            "Severity",
            "Gets a label row",
            "Rule note",
            "Written to match",
            "Written not to match",
            "Category",
        ],
        body,
        st,
        widths=[11, 18, 8, 12, 50, 36, 30, 40],
    )


def _wilson(p: str, n: str, sign: str) -> str:
    edge = "MAX(0," if sign == "-" else "MIN(1,"
    return f'=IF({p}="","",{edge}({p}+Z95^2/(2*{n}){sign}Z95*SQRT({p}*(1-{p})/{n}+Z95^2/(4*{n}^2)))/(1+Z95^2/{n})))'


def _write_metrics(
    ws, kinds, ranges, rules, categories, reasons, guide, sample_info, rater, round_, st
) -> None:
    from openpyxl.utils import get_column_letter

    for c, w in enumerate([30, 14, 12, 12, 13, 14, 11, 11, 14, 11, 11, 10, 12, 12, 14], start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    pop = int(((sample_info or {}).get("population") or {}).get("no_signal_contents", 0))
    _write_lines(
        ws,
        1,
        [
            ("title", f"Metrics: rater {rater}, round {round_}"),
            (
                "muted",
                "Calculated from this workbook with the same definitions as annotation_kit.py score. "
                "Provisional: the official numbers come from score after you import.",
            ),
            ("font", ""),
            ("h2", "Settings"),
        ],
        st,
    )
    settings = [
        ("z for Wilson 95% intervals", 1.96, "validation.wilson_ci uses z = 1.96"),
        (
            "No-signal contents in the sample",
            pop,
            "SAMPLE_MANIFEST.json population.no_signal_contents; scales the miss rate",
        ),
        ("Strict precision H1 threshold", 0.8, "RESEARCH_QUESTION.md H1"),
    ]
    for i, (k, val, src) in enumerate(settings, start=5):
        _text(ws.cell(row=i, column=1), k, st["font"])
        cell = ws.cell(row=i, column=2, value=val)
        cell.font = st["font"]
        _text(ws.cell(row=i, column=3), src, st["muted"])
    # B5 is the defined name Z95; B6 the population; B7 the H1 threshold.
    row = 9
    _text(ws.cell(row=row, column=1), "Progress and pace", st["h2"])
    row += 1
    head = [
        "Kind",
        "Rows",
        "Done",
        "Share done",
        "To do",
        "Need a reason",
        "Invalid",
        "Minutes logged",
        "Avg min per row",
        "Due",
        "Days left",
        "Rows a day needed",
        "Hours left, your pace",
        "Hours left, guideline",
        "Round 2 from (all)",
    ]
    for c, h in enumerate(head, start=1):
        cell = ws.cell(row=row, column=c)
        _text(cell, h, st["head"])
        cell.fill = st["headfill"]
        cell.alignment = st["wrap"]
    ws.row_dimensions[row].height = 30
    lock = {}
    for kind in [k for k in v.KINDS if k in kinds]:
        row += 1
        R = ranges[kind]
        sh = R["sheet"]
        key_col = "file_sha" if kind == "signals" else "file_sha_a"
        rr = R["rng"]
        per_row = _guideline_minutes_per_row(kind, kinds[kind]["items"], guide)
        due = label_ui.ROUND_DUE.get((kind, round_), "")
        gap = label_ui.ROUND2_MIN_DAYS[kind]
        cells = [
            sh,
            f"=COUNTA({sh}!{rr[key_col]})",
            f'=COUNTIF({sh}!{rr["status"]},"done")',
            f'=IF(B{row}=0,"",C{row}/B{row})',
            f"=B{row}-C{row}",
            f'=COUNTIF({sh}!{rr["status"]},"needs a reason")+COUNTIF({sh}!{rr["status"]},"needs a category")',
            f'=COUNTIF({sh}!{rr["status"]},"invalid label")',
            f"=SUM({sh}!{rr['minutes']})",
            f'=IF(COUNT({sh}!{rr["minutes"]})=0,"",H{row}/COUNT({sh}!{rr["minutes"]}))',
            date.fromisoformat(due) if due else None,
            f'=IF(J{row}="","",MAX(0,J{row}-TODAY()+1))',
            f'=IF(E{row}=0,0,IF(OR(K{row}="",K{row}=0),E{row},ROUNDUP(E{row}/K{row},0)))',
            f'=IF(I{row}="","",ROUND(E{row}*I{row}/60,1))',
            f"=ROUND(E{row}*{per_row}/60,1)",
            f'=IF(OR(C{row}<B{row},COUNT({sh}!{rr["date"]})=0),"",MAX({sh}!{rr["date"]})+{gap})',
        ]
        for c, val in enumerate(cells, start=1):
            cell = ws.cell(row=row, column=c)
            if c == 1:
                _text(cell, val, st["bold"])
            else:
                cell.value = val
                cell.font = st["font"]
                cell.fill = st["calc"]
            if c == 4:
                cell.number_format = "0%"
            if c in (10, 15):
                cell.number_format = "yyyy-mm-dd"
        lock[kind] = f"AND($B${row}>0,$C${row}=$B${row})"
    row += 1
    _text(
        ws.cell(row=row, column=1),
        "Hours left, guideline: the guideline's time per item (section 7) spread over the rows. Round 2 from (all): the latest Date labelled plus the gap, once every row is done.",
        st["muted"],
    )
    row += 2
    _text(
        ws.cell(row=row, column=1),
        "Results (provisional; each table unlocks when every row of its kind is done)",
        st["h2"],
    )
    row += 1
    _text(
        ws.cell(row=row, column=1),
        "Locked while you label, because seeing running results can nudge the labels that follow.",
        st["muted"],
    )
    row += 2

    def header(cols_):
        nonlocal row
        for c, h in enumerate(cols_, start=1):
            cell = ws.cell(row=row, column=c)
            _text(cell, h, st["head"])
            cell.fill = st["headfill"]
            cell.alignment = st["wrap"]
        ws.row_dimensions[row].height = 30
        row += 1

    def put(r, c, val, fmt=None):
        cell = ws.cell(row=r, column=c)
        cell.value = val
        cell.font = st["font"]
        cell.fill = st["calc"]
        if fmt:
            cell.number_format = fmt

    if "signals" in kinds:
        R = ranges["signals"]["rng"]
        lk = lock["signals"]
        sh = "Signals"
        _text(
            ws.cell(row=row, column=1),
            "Signals: precision by category (counted once per skill)",
            st["bold"],
        )
        put(row, 2, f'=IF({lk},"unlocked","locked")')
        row += 1
        header(
            [
                "Category",
                "n",
                "RISKY",
                "BENIGN_CONTEXT",
                "NOT_PRESENT",
                "Strict precision",
                "CI low",
                "CI high",
                "Capability precision",
                "CI low",
                "CI high",
                "Meets H1",
            ]
        )
        max_sev: dict[str, int] = {}
        for r in rules.values():
            max_sev[str(r["category"])] = max(
                max_sev.get(str(r["category"]), 0), int(r["severity"])
            )
        cats = sorted(c for c, s in max_sev.items() if s >= label_ui.HIGH_RISK_SEVERITY)
        first = row
        for cat in cats:
            grp = f"{sh}!{R['category']},$A{row},{sh}!{R['h_first']},1,{sh}!{R['h_cat']}"
            _text(ws.cell(row=row, column=1), cat, st["font"])
            put(row, 2, f'=IF({lk},COUNTIFS({grp},">=0"),"")')
            put(row, 3, f'=IF({lk},COUNTIFS({grp},2),"")')
            put(row, 4, f'=IF({lk},COUNTIFS({grp},1),"")')
            put(row, 5, f'=IF({lk},COUNTIFS({grp},0),"")')
            put(row, 6, f'=IF(OR(B{row}="",B{row}=0),"",C{row}/B{row})', "0.00")
            put(row, 7, _wilson(f"F{row}", f"B{row}", "-"), "0.00")
            put(row, 8, _wilson(f"F{row}", f"B{row}", "+"), "0.00")
            put(row, 9, f'=IF(OR(B{row}="",B{row}=0),"",(C{row}+D{row})/B{row})', "0.00")
            put(row, 10, _wilson(f"I{row}", f"B{row}", "-"), "0.00")
            put(row, 11, _wilson(f"I{row}", f"B{row}", "+"), "0.00")
            put(row, 12, f'=IF(F{row}="","",IF(F{row}>=$B$7,"yes","no"))')
            row += 1
        put(
            row,
            1,
            f'=IF(NOT({lk}),"",IF(COUNTIF(L{first}:L{row - 1},"yes")>0,"H1 met for at least one category (provisional)","H1 not met by any category (provisional)"))',
        )
        row += 2

        _text(ws.cell(row=row, column=1), "Signals: precision by rule", st["bold"])
        row += 1
        header(
            [
                "Rule",
                "Category",
                "n",
                "RISKY",
                "BENIGN_CONTEXT",
                "NOT_PRESENT",
                "Strict precision",
                "CI low",
                "CI high",
            ]
        )
        for rid, r in rules.items():
            if int(r["severity"]) < label_ui.HIGH_RISK_SEVERITY:
                continue
            grp = f"{sh}!{R['rule_id']},$A{row},{sh}!{R['h_rank']}"
            _text(ws.cell(row=row, column=1), rid, st["font"])
            _text(ws.cell(row=row, column=2), str(r["category"]), st["font"])
            put(row, 3, f'=IF({lk},COUNTIFS({grp},">=0"),"")')
            put(row, 4, f'=IF({lk},COUNTIFS({grp},2),"")')
            put(row, 5, f'=IF({lk},COUNTIFS({grp},1),"")')
            put(row, 6, f'=IF({lk},COUNTIFS({grp},0),"")')
            put(row, 7, f'=IF(OR(C{row}="",C{row}=0),"",D{row}/C{row})', "0.00")
            put(row, 8, _wilson(f"G{row}", f"C{row}", "-"), "0.00")
            put(row, 9, _wilson(f"G{row}", f"C{row}", "+"), "0.00")
            row += 1
        row += 1

        _text(
            ws.cell(row=row, column=1),
            "Signals: what settled the false positives (NOT_PRESENT and BENIGN_CONTEXT), by rule",
            st["bold"],
        )
        row += 1
        ids = [r["id"] for r in reasons]
        header(["Rule"] + ids + ["not given"])
        for rid, r in rules.items():
            if int(r["severity"]) < label_ui.HIGH_RISK_SEVERITY:
                continue
            base = f'{sh}!{R["rule_id"]},$A{row},{sh}!{R["h_rank"]},"<2"'
            _text(ws.cell(row=row, column=1), rid, st["font"])
            for c in range(2, len(reasons) + 2):
                put(row, c, f'=IF({lk},COUNTIFS({base},{sh}!{R["why"]},Lists!$D${c}),"")')
            put(row, len(reasons) + 2, f'=IF({lk},COUNTIFS({base},{sh}!{R["why"]},""),"")')
            row += 1
        row = _write_lines(ws, row, [("muted", r["text"]) for r in reasons], st) + 1

        _text(
            ws.cell(row=row, column=1),
            "Signals: misses among the skills no rule matched",
            st["bold"],
        )
        row += 1
        header(
            [
                "Scope",
                "n labelled",
                "Missed",
                "Miss rate",
                "CI low",
                "CI high",
                "Est. missed contents",
                "Est. low",
                "Est. high",
            ]
        )
        none_base = f'{sh}!{R["rule_id"]},"NONE",{sh}!{R["status"]},"done"'
        scopes = [("ALL", "")] + [(c, c) for c in categories]
        for scope, cat in scopes:
            _text(ws.cell(row=row, column=1), scope, st["font"])
            put(row, 2, f'=IF({lk},COUNTIFS({none_base}),"")')
            missed = f'COUNTIFS({none_base},{sh}!{R["label"]},"MISSED_RISKY"' + (
                f",{sh}!{R['missed_category']},$A{row})" if cat else ")"
            )
            put(row, 3, f'=IF({lk},{missed},"")')
            put(row, 4, f'=IF(OR(B{row}="",B{row}=0),"",C{row}/B{row})', "0.000")
            put(row, 5, _wilson(f"D{row}", f"B{row}", "-"), "0.000")
            put(row, 6, _wilson(f"D{row}", f"B{row}", "+"), "0.000")
            put(row, 7, f'=IF(D{row}="","",D{row}*$B$6)', "#,##0")
            put(row, 8, f'=IF(E{row}="","",E{row}*$B$6)', "#,##0")
            put(row, 9, f'=IF(F{row}="","",F{row}*$B$6)', "#,##0")
            row += 1
        row += 1

    if "lineage" in kinds:
        R = ranges["lineage"]["rng"]
        lk = lock["lineage"]
        _text(ws.cell(row=row, column=1), "Lineage precision", st["bold"])
        put(row, 2, f'=IF({lk},"unlocked","locked")')
        row += 1
        header(
            [
                "",
                "YES",
                "NO",
                "UNSURE",
                "Precision (YES of YES or NO)",
                "CI low",
                "CI high",
                "Share UNSURE",
            ]
        )
        _text(ws.cell(row=row, column=1), "Lineage", st["font"])
        for c, lab in enumerate(("YES", "NO", "UNSURE"), start=2):
            put(
                row,
                c,
                f'=IF({lk},COUNTIFS(Lineage!{R["status"]},"done",Lineage!{R["label"]},"{lab}"),"")',
            )
        put(row, 5, f'=IF(B{row}="","",IF(B{row}+C{row}=0,"",B{row}/(B{row}+C{row})))', "0.00")
        put(row, 6, _wilson(f"E{row}", f"(B{row}+C{row})", "-"), "0.00")
        put(row, 7, _wilson(f"E{row}", f"(B{row}+C{row})", "+"), "0.00")
        put(
            row,
            8,
            f'=IF(B{row}="","",IF(B{row}+C{row}+D{row}=0,"",D{row}/(B{row}+C{row}+D{row})))',
            "0.00",
        )
        row += 3

    if "drift" in kinds:
        R = ranges["drift"]["rng"]
        lk = lock["drift"]
        _text(ws.cell(row=row, column=1), "Drift: change types", st["bold"])
        put(row, 2, f'=IF({lk},"unlocked","locked")')
        row += 1
        header(["Change type", "Count", "Share", "CI low", "CI high"])
        first = row
        n_labels = len(v.DRIFT_LABELS)
        total = f"SUM($B${first}:$B${first + n_labels - 1})"
        for lab in v.DRIFT_LABELS:
            _text(ws.cell(row=row, column=1), lab, st["font"])
            put(
                row,
                2,
                f'=IF({lk},COUNTIFS(Drift!{R["status"]},"done",Drift!{R["label"]},$A{row}),"")',
            )
            put(row, 3, f'=IF(OR(B{row}="",{total}=0),"",B{row}/{total})', "0.00")
            put(row, 4, _wilson(f"C{row}", total, "-"), "0.00")
            put(row, 5, _wilson(f"C{row}", total, "+"), "0.00")
            row += 1
        row += 1

    _text(ws.cell(row=row, column=1), "Time", st["bold"])
    row += 1
    first_kind = 11
    last_kind = 10 + len([k for k in v.KINDS if k in kinds])
    _text(ws.cell(row=row, column=1), "Hours logged, all kinds", st["font"])
    put(row, 2, f"=ROUND(SUM(H{first_kind}:H{last_kind})/60,1)", "0.0")
    row += 1
    _text(
        ws.cell(row=row, column=1),
        "Definitions: strict precision = RISKY / n; capability precision = (RISKY + BENIGN_CONTEXT) / n; a skill's label for a "
        "category is the highest of its rule rows in that category (RISKY > BENIGN_CONTEXT > NOT_PRESENT); miss rate = MISSED_RISKY / "
        "no-signal rows labelled; lineage precision = YES / (YES + NO); intervals are Wilson 95%.",
        st["muted"],
    )


def _write_assignment_map(ws, kinds, st) -> None:
    body = []
    for kind in kinds:
        for metric, answers, goes in label_ui.ASSIGNMENT_USE[kind]:
            body.append([kind, metric, answers, goes])
    row = _write_lines(
        ws,
        1,
        [
            ("title", "Assignment map: what each number is for"),
            (
                "muted",
                "Rubric: Sprint 2 asks for a validation sample and annotation protocol and preliminary results; Sprint 3 for error analysis and manually inspected examples, and a threats section that covers measurement error.",
            ),
            ("font", ""),
        ],
        st,
    )
    _write_table(
        ws,
        row,
        ["Kind", "Number", "What it answers", "Where it goes"],
        body,
        st,
        widths=[10, 46, 56, 64],
    )


# ---------------------------------------------------------------------------
# Reading a filled-in workbook
# ---------------------------------------------------------------------------


def _get(row, index: dict[str, int], key: str):
    return row[index[key]].value


def _cell_text(val) -> str:
    if val is None:
        return ""
    if isinstance(val, float) and val.is_integer():
        return str(int(val))
    return str(val).strip()


def read_workbook(path: Path) -> dict:
    """Labels and work log from a workbook this module wrote.

    Returns {"rater", "round", "labels": {kind: DataFrame in the label-file format},
    "worklog": {kind: DataFrame in WORKLOG_COLUMNS}}. Raises ValueError for any other file.
    """
    _require_openpyxl()
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=False)
    if "_meta" not in wb.sheetnames:
        raise ValueError(f"{path.name}: not a labelling workbook from annotation_kit.py xlsx")
    meta = {
        str(r[0].value): "" if r[1].value is None else str(r[1].value)
        for r in wb["_meta"].iter_rows(min_row=1, max_col=2)
        if r[0].value
    }
    if meta.get("format") != FORMAT:
        raise ValueError(f"{path.name}: unknown workbook format {meta.get('format')!r}")
    rater, round_ = meta["rater"], int(meta["round"])
    out = {"rater": rater, "round": round_, "labels": {}, "worklog": {}}
    for kind in [k for k in meta.get("kinds", "").split(",") if k]:
        ws = wb[SHEET[kind]]
        headers = {cell.value: i for i, cell in enumerate(ws[1])}
        index = {key: headers[header] for key, header, _w, _r in COLUMNS[kind]}
        label_rows, log_rows = [], []
        for row in ws.iter_rows(min_row=2):
            if kind == "signals":
                key_a, key_b = _cell_text(_get(row, index, "file_sha")), ""
                if not key_a:
                    continue
                rid = _cell_text(_get(row, index, "rule_id"))
                label_rows.append(
                    {
                        "file_sha": key_a,
                        "rule_id": rid,
                        "label": _cell_text(_get(row, index, "label")),
                        "missed_category": _cell_text(_get(row, index, "missed_category")),
                        "notes": _cell_text(_get(row, index, "notes")),
                    }
                )
                prefix = _cell_text(_get(row, index, "why")).split(":", 1)[0]
                why = prefix if re.match(r"^(Q\d+|S|O)$", prefix) else ""
                item_no = _cell_text(_get(row, index, "item"))
            else:
                key_a, key_b = (
                    _cell_text(_get(row, index, "file_sha_a")),
                    _cell_text(_get(row, index, "file_sha_b")),
                )
                if not key_a:
                    continue
                rid, why = "", ""
                label_rows.append(
                    {
                        "file_sha_a": key_a,
                        "file_sha_b": key_b,
                        v.KIND_LABEL_COLUMN[kind]: _cell_text(_get(row, index, "label")),
                        "notes": _cell_text(_get(row, index, "notes")),
                    }
                )
                item_no = _cell_text(_get(row, index, "row"))
            d = _get(row, index, "date")
            first = (
                d.date().isoformat()
                if isinstance(d, datetime)
                else (d.isoformat() if isinstance(d, date) else _cell_text(d))
            )
            minutes = pd.to_numeric(_get(row, index, "minutes"), errors="coerce")
            conf = _cell_text(_get(row, index, "confidence"))
            log_rows.append(
                {
                    "kind": kind,
                    "rater": rater,
                    "round": str(round_),
                    "item_no": item_no,
                    "key_a": key_a,
                    "key_b": key_b,
                    "rule_id": rid,
                    "first_labelled_at": first,
                    "last_changed_at": "",
                    "label_changes": "",
                    "confidence": conf if conf in label_ui.CONFIDENCE else "",
                    "reason_type": why,
                    "active_seconds": ""
                    if pd.isna(minutes)
                    else str(round(float(minutes) * 60, 1)),
                }
            )
        out["labels"][kind] = pd.DataFrame(label_rows, columns=v.KIND_COLUMNS[kind]).fillna("")
        out["worklog"][kind] = pd.DataFrame(log_rows, columns=label_ui.WORKLOG_COLUMNS).fillna("")
    return out


def merge_worklog(new: pd.DataFrame, old: pd.DataFrame | None) -> pd.DataFrame:
    """Fill the workbook's work log from an existing one where the workbook left a field empty."""
    if old is None or old.empty:
        return new
    keys = ["key_a", "key_b", "rule_id"]
    old = old.astype(str).set_index(keys)
    merged = new.astype(str).set_index(keys)
    for col in label_ui.WORKLOG_COLUMNS:
        if col in keys or col not in old.columns:
            continue
        fill = old[col].reindex(merged.index)
        empty = merged[col].str.strip() == ""
        merged.loc[empty, col] = fill[empty].fillna("")
    return merged.reset_index()[label_ui.WORKLOG_COLUMNS]
